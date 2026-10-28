"""Eines per als tests: un GTFS petit fet a mida (dades de prova, no del joc)."""

import io
import pathlib
import sys
import zipfile

ARREL = pathlib.Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ARREL / "scripts"))

GTFS_PROVA = {
    "agency.txt": "agency_id,agency_name,agency_url,agency_timezone\nA,Operador de prova,https://exemple.cat,Europe/Madrid\n",
    "stops.txt": (
        "stop_id,stop_name,stop_lat,stop_lon,location_type,parent_station\n"
        "E1,Estació Prova,41.40000,2.20000,1,\n"
        "E1a,Estació Prova,41.40010,2.20000,0,E1\n"
        "E1b,Estació Prova,41.39990,2.20000,0,E1\n"
        "B1,Pl. de Prova,41.40100,2.20100,0,\n"
        "B2,Plaça Prova,41.40110,2.20110,0,\n"
        "B3,Lluny,41.50000,2.30000,0,\n"
    ),
    "routes.txt": (
        "route_id,agency_id,route_short_name,route_long_name,route_type,route_color,route_text_color\n"
        "M1,A,L1,Línia 1,1,CE1126,FFFFFF\n"
        "H1,A,H12,Bus H12,3,,\n"
    ),
    "calendar.txt": (
        "service_id,monday,tuesday,wednesday,thursday,friday,saturday,sunday,start_date,end_date\n"
        "FEINER,1,1,1,1,1,0,0,20261102,20261231\n"
        "CAPSETMANA,0,0,0,0,0,1,1,20261102,20261231\n"
    ),
    # El primer dimecres (4 de novembre) no hi ha servei de feiner: el dia de referència ha de ser l'11.
    "calendar_dates.txt": "service_id,date,exception_type\nFEINER,20261104,2\n",
    "trips.txt": (
        "route_id,service_id,trip_id\n"
        "M1,FEINER,m1\nM1,FEINER,m2\nM1,CAPSETMANA,m3\n"
        "H1,FEINER,h1\nH1,FEINER,h2\nH1,FEINER,hf\n"
    ),
    # hf funciona per freqüències: de 08:00 a 09:00 cada 10 minuts són 6 expedicions.
    "frequencies.txt": "trip_id,start_time,end_time,headway_secs\nhf,08:00:00,09:00:00,600\n",
    "stop_times.txt": (
        "trip_id,arrival_time,departure_time,stop_id,stop_sequence\n"
        "m1,08:00:00,08:00:00,E1a,1\nm2,09:00:00,09:00:00,E1b,1\nm3,10:00:00,10:00:00,E1a,1\n"
        "h1,08:00:00,08:00:00,B1,1\nh1,08:05:00,08:05:00,B3,2\n"
        "h2,09:00:00,09:00:00,B2,1\n"
        "hf,08:00:00,08:00:00,B1,1\n"
    ),
}


def crea_gtfs(cami, fitxers=None):
    with zipfile.ZipFile(cami, "w") as z:
        for nom, contingut in (fitxers or GTFS_PROVA).items():
            z.writestr(nom, contingut)
    return cami


def llibre(files_per_full):
    """Un llibre d'openpyxl en memòria: {nom del full: [files]}."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nom, files in files_per_full.items():
        full = wb.create_sheet(nom)
        for fila in files:
            full.append(fila)
    dades = io.BytesIO()
    wb.save(dades)
    dades.seek(0)
    return openpyxl.load_workbook(dades, read_only=True, data_only=True)
