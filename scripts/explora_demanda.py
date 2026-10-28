#!/usr/bin/env python3
"""Exploració temporal de les fonts de demanda (fulls de càlcul de TMB i de Renfe, pàgines de l'OMC)."""

import io
import pathlib
import re
import urllib.request

import openpyxl

NAVEGADOR = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "ca,es;q=0.8,en;q=0.5",
}


def baixa(url):
    return urllib.request.urlopen(urllib.request.Request(url, headers=NAVEGADOR), timeout=120).read()


sortida = pathlib.Path("dades/generades/exploracio")
sortida.mkdir(parents=True, exist_ok=True)
for f in sortida.glob("*"):
    f.unlink()

TMB = "https://www.tmb.cat"
fulls = {
    "fmb_2025": TMB + "/documents/525003/1477764/Resum+dades+mensuals+i+di%C3%A0ries+de+viatgers+FMB+2025.xlsx/5d0569be-2352-3003-c90d-a20de72cd94a",
    "tb_2025": TMB + "/documents/525003/1477766/Resum+dades+mensuals+i+di%C3%A0ries+de+viatgers+TB+2025.xlsx/eeb08133-9af2-a3ac-d850-42d54301a2f2",
    "renfe_2018": "https://data.renfe.com/dataset/9190f983-e138-42da-901a-b37205562fe4/resource/10c14d86-c031-4c99-b8e3-bbc8994b6024/download/barcelona_viajeros_por_franja_xlsx.xlsx",
}
for nom, url in fulls.items():
    linies = [url]
    try:
        llibre = openpyxl.load_workbook(io.BytesIO(baixa(url)), read_only=True, data_only=True)
        for full in llibre.worksheets:
            linies.append(f"=== full «{full.title}» ({full.max_row} files × {full.max_column} columnes)")
            for i, fila in enumerate(full.iter_rows(values_only=True)):
                if i >= 45:
                    break
                linies.append(" | ".join("" if v is None else str(v) for v in fila))
    except Exception as e:  # noqa: BLE001
        linies.append(f"ERROR {e}")
    (sortida / f"{nom}.txt").write_text("\n".join(linies) + "\n")

linies = []
for url in [
    "https://transparencia.tmb.cat/ca/linies-actuacio/dades-basiques-i-demanda",
    "https://omc.cat/ca/w/demanda-a-les-estacions",
    "https://omc.cat/en/w/demand-per-line-and-railway-operator",
]:
    linies.append(f"=== {url}")
    try:
        html = baixa(url).decode("utf-8", "replace")
    except Exception as e:  # noqa: BLE001
        linies.append(f"ERROR {e}")
        continue
    linies.append(f"{len(html)} caràcters")
    for m in re.finditer(r"[\"']([^\"'\s]*(?:json|csv|xlsx|download|documents/)[^\"'\s]*)[\"']", html, re.I):
        linies.append("  " + m.group(1)[:300])
(sortida / "enllacos.txt").write_text("\n".join(linies) + "\n")
